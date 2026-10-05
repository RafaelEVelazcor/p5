# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches, Pt, RGBColor


def blue_run(paragraph, text, size=14, bold=False):
    run = paragraph.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(15, 95, 168)
    run.bold = bold
    return run


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

# Primera página
p = doc.add_paragraph()
blue_run(p, 'Reporte de la Práctica 7 -\n', 24, True)
blue_run(p, 'Pipeline con\n', 24, True)
blue_run(p, 'UAT', 24, True)

p = doc.add_paragraph('Universidad / Institución: ITESO')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('Materia: Desarrollo de Software / Integración y despliegue')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('Nombre del proyecto: Cloudflare Worker con D1 y GitHub Actions')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('URL del proyecto:')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('https://p5.rafitavelazco.workers.dev')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('Fecha: 04 de octubre de 2026')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph()
blue_run(doc.add_paragraph(), '1. Objetivo', 16, True)
p = doc.add_paragraph('El objetivo de esta práctica fue automatizar la validación y despliegue del proyecto mediante una pipeline en GitHub Actions, integrando pruebas unitarias, cobertura de código y publicación en Cloudflare Workers.')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph()
blue_run(doc.add_paragraph(), '2. Descripción del sistema', 16, True)
p = doc.add_paragraph('Se desarrolló un Worker de Cloudflare con acceso a una base de datos D1 para consultar registros de personas mediante el endpoint /api/personas. El proyecto incluye configuración de Wrangler, pruebas con Vitest y despliegue automatizado.')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph()
blue_run(doc.add_paragraph(), '3. Pipeline implementada', 16, True)
p = doc.add_paragraph('Se configuró una workflow con dos jobs:')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('- Build and test: instala dependencias, compila el proyecto, ejecuta pruebas y genera cobertura.')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('- Deploy: publica la aplicación en Cloudflare Workers cuando se ejecuta desde la rama principal.')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('Espacio para imagen 1: captura del workflow de GitHub Actions')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('[Inserte aquí la captura de la pipeline / jobs ejecutados]')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

doc.add_page_break()

# Segunda página
p = doc.add_paragraph()
blue_run(doc.add_paragraph(), '4. Resultados obtenidos', 16, True)
p = doc.add_paragraph('La validación local del proyecto fue exitosa:')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('- npm run build: ejecutado correctamente.')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('- npm run test:coverage: ejecutado correctamente con pruebas pasando.')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('- Cobertura generada con resultado satisfactorio para la práctica.')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('Espacio para imagen 2: evidencia de cobertura o resultados de pruebas')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('[Inserte aquí la captura del reporte de cobertura o salida de Vitest]')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph()
blue_run(doc.add_paragraph(), '5. Evidencia de despliegue', 16, True)
p = doc.add_paragraph('Proyecto desplegado: https://p5.rafitavelazco.workers.dev')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('Endpoint funcional: /api/personas')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('Espacio para imagen 3: captura de la URL publicada')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('[Inserte aquí la captura de la plataforma de Cloudflare y la URL publicada]')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph()
blue_run(doc.add_paragraph(), '6. Evaluación UAT', 16, True)
p = doc.add_paragraph('Se verificó que la aplicación cumple con los requisitos de despliegue y validación en la nube. La funcionalidad principal respondió correctamente, la pipeline ejecutó tareas de compilación y pruebas, y la publicación quedó lista para uso.')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

rows = [
    ['Criterio', 'Cumplimiento', 'Evidencia'],
    ['YAML / Workflow', '', ''],
    ['GitHub Action', '', ''],
    ['UAT Report', '', ''],
    ['Project URL', '', ''],
    ['Scoring / Resultado', '', ''],
]
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
for i, cell in enumerate(table.rows[0].cells):
    cell.text = rows[0][i]
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'Calibri'; run.font.size = Pt(10); run.bold = True
for row in rows[1:]:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].text = text
        for paragraph in cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.name = 'Calibri'; run.font.size = Pt(10)

p = doc.add_paragraph('Observaciones:')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(10); run.bold = True

p = doc.add_paragraph('En esta sección puede incluirse una conclusión breve del instructor o del equipo respecto a la calidad del despliegue, cobertura y automatización.')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

p = doc.add_paragraph('Firma / responsable: _____________________________')
for run in p.runs:
    run.font.name = 'Calibri'; run.font.size = Pt(11)

out = r'C:\Users\rafit\OneDrive - ITESO\p5\Reporte_Practica_7_Pipeline_UAT_2Hojas.docx'
doc.save(out)
print('CREATED', out)
