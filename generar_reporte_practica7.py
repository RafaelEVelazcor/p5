# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def azul(run, size=14, bold=False):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(15, 95, 168)
    run.bold = bold


def set_text(paragraph, text):
    p = paragraph
    p.text = text
    for run in p.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
    if not p.runs:
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(11.69)
section.page_height = Inches(8.27)
section.left_margin = Inches(0.35)
section.right_margin = Inches(0.35)
section.top_margin = Inches(0.4)
section.bottom_margin = Inches(0.35)
section.gutter = Inches(0.2)
sectPr = section._sectPr
cols = sectPr.xpath('./w:cols')
if not cols:
    cols = OxmlElement('w:cols')
    cols.set(qn('w:num'), '3')
    cols.set(qn('w:space'), '420')
    sectPr.append(cols)
else:
    cols[0].set(qn('w:num'), '3')
    cols[0].set(qn('w:space'), '420')

# Columna izquierda
p = doc.add_paragraph()
p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
run = p.add_run('Reporte de la Práctica 7 -\nPipeline con\nUAT')
azul(run, 20, True)

p = doc.add_paragraph('Universidad / Institución: ITESO')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('Materia: Desarrollo de Software / Integración y despliegue')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('Nombre del proyecto: Cloudflare Worker con D1 y GitHub Actions')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('URL del proyecto:')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('https://p5.rafitavelazco.workers.dev')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('Fecha: 04 de octubre de 2026')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph()
p = doc.add_paragraph()
run = p.add_run('1. Objetivo')
azul(run, 14, True)

p = doc.add_paragraph('El objetivo de esta práctica fue automatizar la validación y despliegue del proyecto mediante una pipeline en GitHub Actions, integrando pruebas unitarias, cobertura de código y publicación en Cloudflare Workers.')

p = doc.add_paragraph()
run = p.add_run('2. Descripción del sistema')
azul(run, 14, True)

p = doc.add_paragraph('Se desarrolló un Worker de Cloudflare con acceso a una base de datos D1 para consultar registros de personas mediante el endpoint /api/personas. El proyecto incluye configuración de Wrangler, pruebas con Vitest y despliegue automatizado.')

# Centro
p = doc.add_paragraph()
run = p.add_run('3. Pipeline implementada')
azul(run, 14, True)

p = doc.add_paragraph('Se configuró una workflow con dos jobs:')

p = doc.add_paragraph('- Build and test: instala dependencias, compila el proyecto, ejecuta pruebas y genera cobertura.')

p = doc.add_paragraph('- Deploy: publica la aplicación en Cloudflare Workers cuando se ejecuta desde la rama principal.')

p = doc.add_paragraph('Espacio para imagen 1: captura del workflow de GitHub Actions')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('[Inserte aquí la captura de la pipeline / jobs ejecutados]')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('Espacio para imagen 2: evidencia de cobertura o resultados de pruebas')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('[Inserte aquí la captura del reporte de cobertura o salida de Vitest]')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph()
run = p.add_run('4. Resultados obtenidos')
azul(run, 14, True)

p = doc.add_paragraph('La validación local del proyecto fue exitosa:')

p = doc.add_paragraph('- npm run build: ejecutado correctamente.')

p = doc.add_paragraph('- npm run test:coverage: ejecutado correctamente con pruebas pasando.')

p = doc.add_paragraph('- Cobertura generada con resultado satisfactorio para la práctica.')

p = doc.add_paragraph()
run = p.add_run('5. Evidencia de despliegue')
azul(run, 14, True)

# derecha
p = doc.add_paragraph('Proyecto desplegado:')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('https://p5.rafitavelazco.workers.dev')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('Endpoint funcional: /api/personas')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('Espacio para imagen 3: captura de la URL publicada')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph('[Inserte aquí la captura de la plataforma de Cloudflare y la URL publicada]')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(11)

p = doc.add_paragraph()
run = p.add_run('6. Evaluación UAT')
azul(run, 14, True)

p = doc.add_paragraph('Se verificó que la aplicación cumple con los requisitos de despliegue y validación en la nube. La funcionalidad principal respondió correctamente, la pipeline ejecutó tareas de compilación y pruebas, y la publicación quedó lista para uso.')

# Tabla de evaluación
p = doc.add_paragraph()
rows = [
    ['Criterio', 'Cumplimiento', 'Evidencia'],
    ['YAML / Workflow', '', ''],
    ['GitHub Action', '', ''],
    ['UAT Report', '', ''],
    ['Project URL', '', ''],
    ['Scoring / Resultado', '', ''],
]

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
for i, cell in enumerate(table.rows[0].cells):
    cell.text = rows[0][i]
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.bold = True
for row in rows[1:]:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].text = text
        for paragraph in cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.name = 'Calibri'
                run.font.size = Pt(10)

p = doc.add_paragraph('Observaciones:')
p.runs[0].font.name = 'Calibri'
p.runs[0].font.size = Pt(10)
p.runs[0].bold = True

p = doc.add_paragraph('En esta sección puede incluirse una conclusión breve del instructor o del equipo respecto a la calidad del despliegue, cobertura y automatización.')

p = doc.add_paragraph('Firma / responsable: _____________________________')

out = r'C:\Users\rafit\OneDrive - ITESO\p5\Reporte_Practica_7_Pipeline_UAT_final.docx'
doc.save(out)
print('CREATED', out)
